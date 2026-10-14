"""PowerPoint (.pptx) text extraction with slide numbers (plan §5.2).

Each slide gives up to three passages: its own text (``text``), its speaker notes (``notes``)
and the text read from its pictures (``ocr``). ``slide_no`` is the slide's position in the
deck, hidden slides included, so it matches the page of a PDF exported from a copy made by
``unhide_all``.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
import zlib
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError, PythonPptxError
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from studybuddy.ingest import ocr as ocr_engine
from studybuddy.ingest.model import IngestError, PassageDraft, clean_text

if TYPE_CHECKING:
    from pptx.chart.chart import Chart
    from pptx.shapes.base import BaseShape
    from pptx.slide import Slide
    from pptx.table import Table

log = logging.getLogger(__name__)

REPEATED_OCR_SLIDES = 3  # the same picture text on this many slides is a logo or branding

_SKIPPED_PLACEHOLDERS = frozenset({PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE})
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .ppt, or an encrypted .pptx
_NON_WORD = re.compile(r"[^\w]+")
_UNREADABLE = (
    PythonPptxError, KeyError, ValueError, OSError, EOFError, zipfile.BadZipFile, zlib.error, etree.LxmlError
)
_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def _xpath(expr: str) -> etree.XPath:
    return etree.XPath(expr, namespaces=_NS)


_PARAGRAPHS = _xpath(".//a:p")
_PARAGRAPH_PIECES = _xpath(".//a:t/text() | .//m:t/text() | .//a:br")
_CHART_TITLES = _xpath(".//c:title")
_CHART_TITLE_CELL = _xpath("./c:tx/c:strRef//c:v/text()")
_CHART_GROUPS = _xpath("./c:chart/c:plotArea/*[c:ser]")
_CATEGORY_LABELS = _xpath("./c:ser[1]/c:cat//c:pt/c:v/text()")
_SERIES_NAMES = _xpath("./c:ser/c:tx//c:v/text()")
_ALTERNATE_CHOICES = _xpath("./mc:AlternateContent/mc:Choice[1]")
_FIRST_OFFSET = _xpath("(.//a:off)[1]")
_DIAGRAM_DATA_RID = _xpath("./a:graphic/a:graphicData/dgm:relIds/@r:dm")
_DIAGRAM_POINTS = _xpath("./dgm:ptLst/dgm:pt")
_DIAGRAM_POINT_TEXT = _xpath("./dgm:t")
_DIAGRAM_PARENT_LINKS = _xpath("./dgm:cxnLst/dgm:cxn[not(@type) or @type='parOf']")

_Position = tuple[int, int]  # (top, left) in EMU


@dataclass
class _Slide:
    no: int
    text: str
    notes: str
    pictures: list[str]  # OCR text of each picture, in reading order


def extract(path: Path, *, ocr: bool = True) -> list[PassageDraft]:
    """Passages for every slide, in deck order, with 1-based ``slide_no``."""
    path = Path(path)
    use_ocr = ocr and ocr_engine.available()
    ocr_cache: dict[str, str] = {}  # image SHA-1 -> text
    slides = []
    for no, slide in enumerate(_slides(path), start=1):
        pictures = [_picture_text(p, ocr_cache) for p in _pictures(slide.shapes)] if use_ocr else []
        slides.append(_Slide(no, _slide_text(slide), _notes_text(slide), pictures))

    repeated = _repeated_ocr(slides)
    passages = []
    for s in slides:
        if s.text:
            passages.append(PassageDraft(s.text, origin="text", slide_no=s.no))
        if s.notes:
            passages.append(PassageDraft(s.notes, origin="notes", slide_no=s.no))
        if ocr_text := _ocr_text(s, repeated):
            passages.append(PassageDraft(ocr_text, origin="ocr", slide_no=s.no))
    log.info("%s: %d slides, %d passages", path.name, len(slides), len(passages))
    return passages


def slide_count(path: Path) -> int:
    return len(_slides(Path(path)))


def hidden_slides(path: Path) -> list[int]:
    """1-based numbers of the slides marked hidden."""
    return [no for no, slide in enumerate(_slides(Path(path)), start=1) if _is_hidden(slide)]


def unhide_all(src: Path, dst: Path) -> int:
    """Write a copy of ``src`` to ``dst`` with every slide visible; return how many were hidden.

    Only the hidden slides' XML is rewritten; every other zip member is copied byte for byte.
    """
    src, dst = Path(src), Path(dst)
    if src.resolve() == dst.resolve():
        raise ValueError(f"unhide_all writes a copy; dst must differ from src ({src})")
    hidden = {slide.part.partname.membername for slide in _slides(src) if _is_hidden(slide)}
    tmp = dst.with_name(dst.name + ".part")
    try:
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w") as zout:
            for info in zin.infolist():
                data = zin.read(info)
                zout.writestr(info, _without_show(data) if info.filename in hidden else data)
        tmp.replace(dst)
    finally:
        tmp.unlink(missing_ok=True)
    return len(hidden)


def diagram_text(xml: bytes) -> str:
    """Node text of a SmartArt data part (``ppt/diagrams/dataN.xml``), in outline order."""
    model = _parse(xml)
    texts: dict[str, str] = {}
    roots = []
    for point in _DIAGRAM_POINTS(model):
        point_id = point.get("modelId")
        texts[point_id] = "\n".join(_text_of(t) for t in _DIAGRAM_POINT_TEXT(point))
        if point.get("type") == "doc":
            roots.append(point_id)
    children: dict[str, list[tuple[int, str]]] = defaultdict(list)
    for link in _DIAGRAM_PARENT_LINKS(model):
        children[link.get("srcId")].append((int(link.get("srcOrd", "0")), link.get("destId")))
    ordered = (texts.get(point_id, "") for point_id in _outline_order(list(texts), roots, children))
    return "\n".join(text for text in ordered if text)


# --- opening -------------------------------------------------------------------------------


def _slides(path: Path) -> list[Slide]:
    if not path.is_file():
        raise IngestError(f"{path}: no such file")
    try:
        return list(Presentation(str(path)).slides)
    except _UNREADABLE as exc:
        raise IngestError(_unreadable(path, exc)) from exc


def _unreadable(path: Path, exc: Exception) -> str:
    try:
        with path.open("rb") as f:
            is_ole = f.read(len(_OLE_MAGIC)) == _OLE_MAGIC
    except OSError:
        is_ole = False
    if is_ole:
        return f"{path.name} is a legacy .ppt or a password-protected file; save it as an unprotected .pptx"
    reason = "not a zip archive" if isinstance(exc, PackageNotFoundError) else f"{type(exc).__name__}: {exc}"
    return f"{path.name} is not a readable .pptx file; it may be damaged or another format ({reason})"


def _is_hidden(slide: Slide) -> bool:
    return slide.element.get("show") in ("0", "false")


def _without_show(slide_xml: bytes) -> bytes:
    root = _parse(slide_xml)
    root.attrib.pop("show", None)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _parse(xml: bytes) -> etree._Element:
    return etree.fromstring(xml, etree.XMLParser(resolve_entities=False, no_network=True))


# --- slide text ----------------------------------------------------------------------------


def _slide_text(slide: Slide) -> str:
    title = slide.shapes.title
    head = [_shape_text(title)] if title is not None else []
    return clean_text("\n".join(head + _blocks(slide.shapes, slide.element.cSld.spTree, skip=title)))


def _notes_text(slide: Slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return clean_text(frame.text.replace("\v", "\n")) if frame is not None else ""


def _blocks(shapes: Iterable[BaseShape], tree: etree._Element, *, skip: BaseShape | None = None) -> list[str]:
    """The text of each shape in reading order (top, then left); a group is one block."""
    found = [(_position(s), _shape_text(s)) for s in shapes if skip is None or s.element is not skip.element]
    found += _alternate_content(tree)
    found.sort(key=lambda item: item[0])
    return [text for _, text in found if text]


def _position(shape: BaseShape) -> _Position:
    return (shape.top or 0, shape.left or 0)


def _shape_text(shape: BaseShape) -> str:
    try:
        if isinstance(shape, GroupShape):
            return "\n".join(_blocks(shape.shapes, shape.element))
        if shape.is_placeholder and shape.placeholder_format.type in _SKIPPED_PLACEHOLDERS:
            return ""
        if shape.has_text_frame:
            return shape.text_frame.text.replace("\v", "\n")
        if isinstance(shape, GraphicFrame):
            if shape.has_table:
                return _table_text(shape.table)
            if shape.has_chart:
                return _chart_text(shape.chart)
            return _smartart_text(shape)
    except Exception as exc:  # one malformed shape must not cost the whole deck
        log.warning("skipped shape %r that could not be read: %s", shape.name, exc)
    return ""


def _alternate_content(tree: etree._Element) -> list[tuple[_Position, str]]:
    """Shapes python-pptx does not see: PowerPoint wraps text boxes holding equations in
    ``mc:AlternateContent``. The first ``mc:Choice`` holds the real text, the fallback a picture."""
    found = []
    for choice in _ALTERNATE_CHOICES(tree):
        offset = _FIRST_OFFSET(choice)
        position = (int(offset[0].get("y", 0)), int(offset[0].get("x", 0))) if offset else (0, 0)
        found.append((position, _text_of(choice)))
    return found


def _text_of(element: etree._Element) -> str:
    """The non-empty paragraphs under ``element``, equations included as their linear text."""
    paragraphs = (
        "".join(piece if isinstance(piece, str) else "\n" for piece in _PARAGRAPH_PIECES(p))
        for p in _PARAGRAPHS(element)
    )
    return "\n".join(p for p in paragraphs if p.strip())


def _table_text(table: Table) -> str:
    rows = (" | ".join(" ".join(c.text.split()) for c in row.cells if not c.is_spanned) for row in table.rows)
    return "\n".join(row for row in rows if row.strip(" |"))


def _chart_text(chart: Chart) -> str:
    """Chart and axis titles, category labels and series names; the plotted numbers are not kept.

    Read from the chart XML, because ``chart.plots`` rejects 3-D, stock and surface charts.
    """
    space = chart.element
    lines = [_text_of(title) or " ".join(_CHART_TITLE_CELL(title)) for title in _CHART_TITLES(space)]
    for group in _CHART_GROUPS(space):  # one per chart type in a combo chart
        lines.append(", ".join(label for label in _CATEGORY_LABELS(group) if label.strip()))
        lines.append(", ".join(name for name in _SERIES_NAMES(group) if name.strip()))
    return "\n".join(dict.fromkeys(line for line in lines if line))


def _smartart_text(frame: GraphicFrame) -> str:
    rids = _DIAGRAM_DATA_RID(frame.element)
    return diagram_text(frame.part.related_part(rids[0]).blob) if rids else ""


def _outline_order(ids: list[str], roots: list[str], children: dict[str, list[tuple[int, str]]]) -> list[str]:
    """Depth-first from the document roots, children by ``srcOrd``; unreached points last."""
    order: list[str] = []
    seen: set[str] = set()
    stack = list(reversed(roots))
    while stack:
        point_id = stack.pop()
        if point_id in seen:
            continue
        seen.add(point_id)
        order.append(point_id)
        stack.extend(child for _, child in sorted(children.get(point_id, ()), reverse=True))
    return order + [point_id for point_id in ids if point_id not in seen]


# --- pictures ------------------------------------------------------------------------------


def _pictures(shapes: Iterable[BaseShape]) -> list[Picture]:
    """Picture shapes (picture placeholders included) in reading order, groups opened."""
    found: list[Picture] = []
    for shape in sorted(shapes, key=_position):
        if isinstance(shape, GroupShape):
            found += _pictures(shape.shapes)
        elif isinstance(shape, Picture):
            found.append(shape)
    return found


def _picture_text(picture: Picture, cache: dict[str, str]) -> str:
    try:
        image = picture.image
    except (KeyError, ValueError, AttributeError) as exc:  # linked, not embedded, or a broken reference
        log.debug("skipped picture %r: %s", picture.name, exc)
        return ""
    if image.sha1 not in cache:
        cache[image.sha1] = _ocr_blob(image.blob)
    return cache[image.sha1]


def _ocr_blob(blob: bytes) -> str:
    try:
        img = PILImage.open(io.BytesIO(blob))
        img.load()
    except (OSError, ValueError, PILImage.DecompressionBombError) as exc:  # EMF/WMF, damaged files
        log.debug("skipped an image Pillow cannot read: %s", exc)
        return ""
    if min(img.size) < ocr_engine.MIN_SIDE_PX:
        return ""
    return clean_text(ocr_engine.ocr_image(img))


def _repeated_ocr(slides: list[_Slide]) -> set[str]:
    found_on: dict[str, set[int]] = defaultdict(set)
    for s in slides:
        for text in s.pictures:
            if text:
                found_on[text].add(s.no)
    repeated = {text for text, nos in found_on.items() if len(nos) >= REPEATED_OCR_SLIDES}
    if repeated:
        log.debug("dropped picture text repeated on %d+ slides: %s", REPEATED_OCR_SLIDES, sorted(repeated))
    return repeated


def _ocr_text(slide: _Slide, repeated: set[str]) -> str:
    """The slide's picture text, minus logos and lines the slide already shows as text."""
    own = f" {_normalize(slide.text)} "
    blocks = []
    for text in dict.fromkeys(slide.pictures):
        if text in repeated:
            continue
        lines = [line for line in text.split("\n") if (n := _normalize(line)) and f" {n} " not in own]
        if lines:
            blocks.append("\n".join(lines))
    return clean_text("\n\n".join(blocks))


def _normalize(text: str) -> str:
    return _NON_WORD.sub(" ", text.casefold()).strip()
