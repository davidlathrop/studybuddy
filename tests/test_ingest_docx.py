"""Word ingestion: heading paths, tables, lists, text boxes, chunking, and placing passages on pages."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml

from doc_samples import PARAGRAPHS, Heading, ListItem, OutlineHeading, Table, TextBox, make_docx, make_text_pdf, wrap
from studybuddy.ingest import docx as word
from studybuddy.ingest import pdf
from studybuddy.ingest.chunking import MAX_WORDS
from studybuddy.ingest.model import IngestError, PassageDraft, word_count

W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _sections(path: Path) -> list[tuple[str | None, str]]:
    return [(d.heading_path, d.text) for d in word.extract(path)]


# --- extract ----------------------------------------------------------------------------------


def test_heading_paths_nest_and_reset(tmp_path: Path) -> None:
    path = make_docx(
        tmp_path / "a.docx",
        [
            "Text before any heading.",
            Heading("Chapter 1", 1),
            "Cells are the unit of life.",
            Heading("1.1 Membranes", 2),
            "Membranes are bilayers.",
            Heading("1.1.1 Channels", 3),
            "Channels pass ions.",
            Heading("1.2 Organelles", 2),
            "Mitochondria make ATP.",
            Heading("Chapter 2", 1),
            "Energy flows through ecosystems.",
        ],
    )
    assert _sections(path) == [
        (None, "Text before any heading."),
        ("Chapter 1", "Chapter 1\n\nCells are the unit of life."),
        ("Chapter 1 > 1.1 Membranes", "1.1 Membranes\n\nMembranes are bilayers."),
        ("Chapter 1 > 1.1 Membranes > 1.1.1 Channels", "1.1.1 Channels\n\nChannels pass ions."),
        ("Chapter 1 > 1.2 Organelles", "1.2 Organelles\n\nMitochondria make ATP."),
        ("Chapter 2", "Chapter 2\n\nEnergy flows through ecosystems."),
    ]
    assert all(d.page_no is None and d.origin == "text" for d in word.extract(path))


def test_title_is_the_top_level_and_bare_headings_join_their_first_subsection(tmp_path: Path) -> None:
    path = make_docx(
        tmp_path / "a.docx",
        [
            Heading("Study Guide", 0),
            Heading("Chapter 2", 1),
            Heading("2.1 Glycolysis", 2),
            "Glucose is split.",
            Heading("Appendix", 1),
            Heading("Glossary", 1),
            "ATP: the energy currency.",
        ],
    )
    assert _sections(path) == [
        (
            "Study Guide > Chapter 2 > 2.1 Glycolysis",
            "Study Guide\n\nChapter 2\n\n2.1 Glycolysis\n\nGlucose is split.",
        ),
        ("Study Guide > Appendix", "Appendix"),  # followed by a sibling, so it stays on its own
        ("Study Guide > Glossary", "Glossary\n\nATP: the energy currency."),
    ]


@pytest.mark.parametrize("via_style", [False, True], ids=["direct", "style"])
def test_outline_level_marks_headings_without_heading_styles(tmp_path: Path, via_style: bool) -> None:
    path = make_docx(
        tmp_path / "a.docx",
        [
            OutlineHeading("Unit A", 1, via_style=via_style),
            "Intro to unit A.",
            OutlineHeading("Topic A1", 2, via_style=via_style),
            "Details of A1.",
            OutlineHeading("Unit B", 1, via_style=via_style),
            "Intro to unit B.",
        ],
    )
    assert [d.heading_path for d in word.extract(path)] == ["Unit A", "Unit A > Topic A1", "Unit B"]


def test_tables_become_one_line_per_row(tmp_path: Path) -> None:
    path = make_docx(
        tmp_path / "a.docx",
        [
            Heading("Respiration", 1),
            Table(
                [["Stage", "Location", "ATP"], ["Glycolysis", "Cytoplasm", "2"], ["", "", ""], ["Krebs cycle", "", "2"]]
            ),
            "After the table.",
        ],
    )
    assert _sections(path) == [
        (
            "Respiration",
            "Respiration\n\nStage | Location | ATP\nGlycolysis | Cytoplasm | 2\nKrebs cycle | 2\n\nAfter the table.",
        )
    ]


def test_list_items_are_marked_and_kept_together(tmp_path: Path) -> None:
    path = make_docx(
        tmp_path / "a.docx",
        [
            "Steps of mitosis:",
            ListItem("Prophase"),
            ListItem("Metaphase", numbered=True),
            ListItem("Anaphase", direct=True),
            "Then cytokinesis.",
            ListItem("Telophase"),
        ],
    )
    assert _sections(path) == [
        (None, "Steps of mitosis:\n\n- Prophase\n- Metaphase\n- Anaphase\n\nThen cytokinesis.\n\n- Telophase")
    ]


def test_text_box_text_follows_its_paragraph_once(tmp_path: Path) -> None:
    path = make_docx(
        tmp_path / "a.docx",
        [
            Heading("Figures", 1),
            TextBox("See the diagram.", ["Box: ATP synthase", "spins like a turbine"]),
            "Next paragraph.",
        ],
    )
    (heading, text), = _sections(path)
    assert text == "Figures\n\nSee the diagram.\n\nBox: ATP synthase\nspins like a turbine\n\nNext paragraph."
    assert text.count("turbine") == 1  # the mc:Fallback copy is skipped


def test_page_headers_and_footers_are_ignored(tmp_path: Path) -> None:
    path = make_docx(tmp_path / "a.docx", ["Only body text."], header="Running Header", footer="Page Footer")
    assert _sections(path) == [(None, "Only body text.")]


def test_tracked_insertions_content_controls_and_toc_are_handled(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph("Kept ")
    p._p.append(parse_xml(f'<w:ins {W_NS} w:id="1" w:author="a"><w:r><w:t>inserted</w:t></w:r></w:ins>'))
    p._p.append(parse_xml(f'<w:del {W_NS} w:id="2" w:author="a"><w:r><w:delText>deleted</w:delText></w:r></w:del>'))
    control = "<w:sdtContent><w:p><w:r><w:t>Inside a content control.</w:t></w:r></w:p></w:sdtContent>"
    doc.element.body.insert(len(doc.element.body) - 1, parse_xml(f"<w:sdt {W_NS}>{control}</w:sdt>"))
    doc.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Chapter 1 ........ 3", style="toc 1")
    doc.save(str(tmp_path / "a.docx"))
    assert _sections(tmp_path / "a.docx") == [(None, "Kept inserted\n\nInside a content control.")]


def test_long_section_is_chunked_under_one_heading_path(tmp_path: Path) -> None:
    paragraphs = list(PARAGRAPHS * 2)
    blocks = [Heading("Unit 3", 1), Heading("Metabolism", 2), *paragraphs, Heading("Unit 4", 1), "Short."]
    path = make_docx(tmp_path / "a.docx", blocks)
    drafts = word.extract(path)
    long = [d for d in drafts if d.heading_path == "Unit 3 > Metabolism"]
    assert len(long) >= 3
    assert long[0].text.startswith("Unit 3\n\nMetabolism\n\n")
    assert all(word_count(d.text) <= MAX_WORDS for d in long)
    assert all(word_count(d.text) >= 150 for d in long), [word_count(d.text) for d in long]
    assert " ".join(d.text for d in long).split() == " ".join(["Unit 3", "Metabolism", *paragraphs]).split()
    assert all(c.text.split("\n\n")[-1] in PARAGRAPHS for c in long), "chunks end at paragraph boundaries"
    assert drafts[-1].heading_path == "Unit 4"


def test_empty_document_has_no_passages(tmp_path: Path) -> None:
    assert word.extract(make_docx(tmp_path / "a.docx", [])) == []


@pytest.mark.parametrize("kind", ["garbage", "zip-without-document", "pdf"])
def test_unreadable_docx_raises_ingest_error(tmp_path: Path, kind: str) -> None:
    path = tmp_path / "bad.docx"
    if kind == "garbage":
        path.write_bytes(b"not a word file")
    elif kind == "zip-without-document":
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("hello.txt", "hi")
    else:
        path.write_bytes(make_text_pdf(tmp_path / "x.pdf", [["PDF"]]).read_bytes())
    with pytest.raises(IngestError, match="not a readable Word document"):
        word.extract(path)


def test_missing_docx_raises_ingest_error(tmp_path: Path) -> None:
    with pytest.raises(IngestError, match="not found"):
        word.extract(tmp_path / "missing.docx")


# --- assign_pages -----------------------------------------------------------------------------


def _drafts(*texts: str) -> list[PassageDraft]:
    return [PassageDraft(t, heading_path=f"H{i}") for i, t in enumerate(texts)]


def _pages(drafts: list[PassageDraft], page_texts: list[str]) -> list[int | None]:
    return [d.page_no for d in word.assign_pages(drafts, page_texts)]


PAGES = [
    "Glycolysis is the first stage of cellular respiration. It takes place in the cytoplasm.",
    "The citric acid cycle runs inside the mitochondrial matrix. Each turn releases carbon dioxide. "
    "Photosynthesis converts light energy into chemical energy stored in sugar.",
    "Enzymes are proteins that speed up chemical reactions by lowering the activation energy.",
]


def test_assign_pages_places_drafts_in_order() -> None:
    drafts = _drafts(
        "Glycolysis is the first stage of cellular respiration. It takes place in the cytoplasm.",
        "The citric acid cycle runs inside the mitochondrial matrix.",
        "Photosynthesis converts light energy into chemical energy stored in sugar.",
        "Enzymes are proteins that speed up chemical reactions.",
    )
    placed = word.assign_pages(drafts, PAGES)
    assert [d.page_no for d in placed] == [1, 2, 2, 3]
    assert [(d.text, d.heading_path) for d in placed] == [(d.text, d.heading_path) for d in drafts]
    assert all(d.page_no is None for d in drafts)  # copies, not modified in place


def test_assign_pages_ignores_hyphenation_ligatures_and_line_breaks() -> None:
    pages = ["Intro page.", "The \ufb01rst stage, glyco-\nlysis, takes place\n\nin the cyto-\nplasm of every cell."]
    drafts = _drafts("Intro page.", "- The first stage, glycolysis, takes place in the cytoplasm of every cell.")
    assert _pages(drafts, pages) == [1, 2]


def test_assign_pages_uses_the_start_page_of_a_spanning_draft() -> None:
    pages = ["Unrelated opening words. Mitosis divides one nucleus", "into two identical nuclei during mitosis."]
    drafts = _drafts("Mitosis divides one nucleus into two identical nuclei during mitosis.")
    assert _pages(drafts, pages) == [1]


def test_assign_pages_falls_back_to_fuzzy_matching_nearby() -> None:
    pages = [
        "Opening page about cells.",
        "Lists: a. Enzymes are b. proteins that speeed up chemcial reactions by lowering the activation energy.",
    ]
    text = "Enzymes are proteins that speed up chemical reactions by lowering the activation energy."
    drafts = _drafts("Opening page about cells.", text)
    assert _pages(drafts, pages) == [1, 2]


def test_assign_pages_never_goes_backwards_and_keeps_the_page_when_lost() -> None:
    drafts = _drafts(
        "Text that is nowhere in the document at all, not even close.",  # before any match: page 1
        "Enzymes are proteins that speed up chemical reactions by lowering the activation energy.",
        "Glycolysis is the first stage of cellular respiration.",  # only on an earlier page
        "Something else that cannot be found anywhere in these pages.",
    )
    assert _pages(drafts, PAGES) == [1, 3, 3, 3]


def test_assign_pages_skips_blank_pages_and_jumps_far_on_an_exact_match() -> None:
    pages = ["Glycolysis is the first stage of cellular respiration.", "", "", "", "", "", "",
             "Photosynthesis converts light energy into chemical energy stored in sugar."]
    drafts = _drafts("Glycolysis is the first stage.", "Photosynthesis converts light energy into chemical energy.")
    assert _pages(drafts, pages) == [1, 8]


def test_assign_pages_prefers_a_near_fuzzy_match_to_a_far_exact_one() -> None:
    text = "Enzymes are proteins that speed up chemical reactions by lowering the activation energy."
    lettered = "a. Enzymes are b. proteins that speed up chemical reactions by lowering the activation energy."
    pages = ["Opening page about cells.", lettered, "", "", "", "", "Review: " + text]
    assert _pages(_drafts("Opening page about cells.", text), pages) == [1, 2]


def test_assign_pages_without_pages_leaves_page_unset() -> None:
    assert _pages(_drafts("Anything at all."), []) == [None]


def test_assign_pages_with_extracted_docx_and_pdf(tmp_path: Path) -> None:
    """Word passages placed on the pages of a PDF holding the same text, as Word would export it."""
    names = "alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron pi".split()
    paragraphs = [f"Note {name}. {text}" for name, text in zip(names, PARAGRAPHS * 2)]
    blocks = [Heading("Chapter One", 1), *paragraphs[:9], Heading("Chapter Two", 1), *paragraphs[9:]]
    docx_path = make_docx(tmp_path / "notes.docx", blocks)
    pages: list[list[str]] = [[]]
    for i, paragraph in enumerate(paragraphs):
        if i == 0:
            pages[-1].append("Chapter One")
        if i == 9:
            pages[-1].append("Chapter Two")
        pages[-1].extend(wrap(paragraph, 70))
        if i % 3 == 2:
            pages.append([])
    pdf_path = make_text_pdf(tmp_path / "notes.pdf", pages, header="Biology notes", footer="Page {page}")
    drafts = word.extract(docx_path)
    assert len(drafts) >= 4
    placed = word.assign_pages(drafts, pdf.page_texts(pdf_path))
    page_of = {line: n for n, lines in enumerate(pages, start=1) for line in lines}
    expected = [page_of[wrap(d.text.split("\n\n")[0], 70)[0]] for d in drafts]
    assert [d.page_no for d in placed] == expected
    assert len(set(expected)) >= 3
