"""Generate small .pptx decks with python-pptx for the PowerPoint ingester's tests.

Like ``media.py``, these helpers never import ``studybuddy``. Shapes that python-pptx cannot
author (SmartArt, equation text boxes, footer placeholders) are injected as XML.
"""

from __future__ import annotations

import copy
import io
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.util import Inches

FONT = "/System/Library/Fonts/Helvetica.ttc"
FONT_SIZE = 48

# Slide numbers in the deck built by make_course_deck().
BULLETS, READING_ORDER, GROUPS, TABLE, CHART, NOTES, HIDDEN = range(1, 8)
FOOTERS, PHOTO, PICTURE_PLACEHOLDER, SMARTART, EQUATION, BLANK = range(8, 14)
SLIDE_COUNT = BLANK

# Layouts of python-pptx's default template.
_TITLE_AND_CONTENT, _TITLE_ONLY, _BLANK, _PICTURE_WITH_CAPTION = 1, 5, 6, 8


@dataclass(frozen=True)
class SampleImage:
    """A PNG with ``lines`` drawn in black; each sample has a distinct ``size`` to tell them apart."""

    lines: tuple[str, ...]
    size: tuple[int, int]
    transparent: bool = False

    @property
    def text(self) -> str:
        return "\n".join(self.lines)


LOGO = SampleImage(("ACME UNIVERSITY",), (560, 110))  # on BULLETS, READING_ORDER and GROUPS
FIGURE = SampleImage(("Membrane transport",), (640, 120))  # in a group on GROUPS, and on TABLE
PHOTO_TEXT = SampleImage(("Light reactions in the thylakoid", "Calvin cycle in the stroma"), (1000, 280))  # PHOTO
CAPTIONED = SampleImage(("Stomata regulate gas exchange",), (900, 130), transparent=True)  # PICTURE_PLACEHOLDER
TINY = SampleImage((), (16, 16))  # on BULLETS; too small to read
IMAGES = (LOGO, FIGURE, PHOTO_TEXT, CAPTIONED, TINY)

PHOTO_SLIDE_LINE = "Calvin cycle in the stroma"  # text box on PHOTO repeating a line of PHOTO_TEXT

NAMESPACES = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)

# A SmartArt data part as PowerPoint writes it, trimmed (model ids may be integers or GUIDs).
# The outline is Interphase (G1 phase, S phase, G2 phase), Mitosis (Prophase), but the points
# are listed out of order, and transition and presentation points carry empty or no text.
SMARTART_OUTLINE = ("Interphase", "G1 phase", "S phase", "G2 phase", "Mitosis", "Prophase")
SMARTART_DATA_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
    xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <dgm:ptLst>
    <dgm:pt modelId="1" type="doc">
      <dgm:prSet loTypeId="urn:microsoft.com/office/officeart/2005/8/layout/hierarchy3" loCatId="list"
          qsTypeId="urn:microsoft.com/office/officeart/2005/8/quickstyle/simple1" qsCatId="simple"
          csTypeId="urn:microsoft.com/office/officeart/2005/8/colors/accent1_2" csCatId="accent1" phldr="1"/>
      <dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="20">
      <dgm:prSet phldrT="[Text]"/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" dirty="0"/><a:t>Mitosis</a:t></a:r></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="12">
      <dgm:prSet phldrT="[Text]"/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/>
        <a:p><a:r><a:rPr lang="en-US"/><a:t>S </a:t></a:r><a:r><a:rPr b="1"/><a:t>phase</a:t></a:r></a:p>
      </dgm:t>
    </dgm:pt>
    <dgm:pt modelId="101" type="parTrans" cxnId="301">
      <dgm:prSet/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="102" type="sibTrans" cxnId="301">
      <dgm:prSet/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="10">
      <dgm:prSet phldrT="[Text]"/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>Interphase</a:t></a:r></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="13">
      <dgm:prSet phldrT="[Text]"/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>G2 phase</a:t></a:r></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="11">
      <dgm:prSet phldrT="[Text]"/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>G1 phase</a:t></a:r></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="21">
      <dgm:prSet phldrT="[Text]"/><dgm:spPr/>
      <dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>Prophase</a:t></a:r></a:p></dgm:t>
    </dgm:pt>
    <dgm:pt modelId="201" type="pres">
      <dgm:prSet presAssocID="1" presName="diagram" presStyleCnt="0">
        <dgm:presLayoutVars><dgm:chPref val="1"/><dgm:dir/><dgm:animOne val="branch"/></dgm:presLayoutVars>
      </dgm:prSet>
      <dgm:spPr/>
    </dgm:pt>
    <dgm:pt modelId="202" type="pres">
      <dgm:prSet presAssocID="10" presName="rootText" presStyleLbl="node0" presStyleIdx="0" presStyleCnt="2"/>
      <dgm:spPr/>
    </dgm:pt>
  </dgm:ptLst>
  <dgm:cxnLst>
    <dgm:cxn modelId="305" srcId="1" destId="20" srcOrd="1" destOrd="0" parTransId="103" sibTransId="104"/>
    <dgm:cxn modelId="303" srcId="10" destId="12" srcOrd="1" destOrd="0" parTransId="105" sibTransId="106"/>
    <dgm:cxn modelId="301" srcId="1" destId="10" srcOrd="0" destOrd="0" parTransId="101" sibTransId="102"/>
    <dgm:cxn modelId="304" srcId="10" destId="13" srcOrd="2" destOrd="0" parTransId="107" sibTransId="108"/>
    <dgm:cxn modelId="302" srcId="10" destId="11" srcOrd="0" destOrd="0" parTransId="109" sibTransId="110"/>
    <dgm:cxn modelId="306" srcId="20" destId="21" srcOrd="0" destOrd="0" parTransId="111" sibTransId="112"/>
    <dgm:cxn modelId="401" type="presOf" srcId="10" destId="202" srcOrd="0" destOrd="0"
        presId="urn:microsoft.com/office/officeart/2005/8/layout/hierarchy3"/>
    <dgm:cxn modelId="402" type="presParOf" srcId="201" destId="202" srcOrd="0" destOrd="0"
        presId="urn:microsoft.com/office/officeart/2005/8/layout/hierarchy3"/>
  </dgm:cxnLst>
  <dgm:bg/>
  <dgm:whole/>
</dgm:dataModel>
"""

_SMARTART_FRAME_XML = f"""<p:graphicFrame {NAMESPACES}>
  <p:nvGraphicFramePr>
    <p:cNvPr id="{{shape_id}}" name="Diagram 1"/><p:cNvGraphicFramePr/><p:nvPr/>
  </p:nvGraphicFramePr>
  <p:xfrm><a:off x="{Inches(1)}" y="{Inches(1.8)}"/><a:ext cx="{Inches(8)}" cy="{Inches(4.5)}"/></p:xfrm>
  <a:graphic>
    <a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">
      <dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
          r:dm="{{rid}}" r:lo="rIdNone1" r:qs="rIdNone2" r:cs="rIdNone3"/>
    </a:graphicData>
  </a:graphic>
</p:graphicFrame>"""

# A text box holding an equation, as PowerPoint saves it: the real shape in mc:Choice
# (Office 2010+), a picture of it with no usable text in mc:Fallback.
EQUATION_TEXT = "Energy is given by E=mc2 in relativity"
_EQUATION_XML = f"""<mc:AlternateContent {NAMESPACES}
    xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
  <mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14">
    <p:sp>
      <p:nvSpPr><p:cNvPr id="{{shape_id}}" name="TextBox 1"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
      <p:spPr>
        <a:xfrm><a:off x="{Inches(1)}" y="{Inches(3)}"/><a:ext cx="{Inches(8)}" cy="{Inches(1)}"/></a:xfrm>
        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      </p:spPr>
      <p:txBody>
        <a:bodyPr wrap="square"/><a:lstStyle/>
        <a:p>
          <a:r><a:rPr lang="en-US"/><a:t>Energy is given by </a:t></a:r>
          <a14:m>
            <m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
              <m:r><m:t>E=m</m:t></m:r>
              <m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>
            </m:oMath>
          </a14:m>
          <a:r><a:rPr lang="en-US"/><a:t> in relativity</a:t></a:r>
        </a:p>
      </p:txBody>
    </p:sp>
  </mc:Choice>
  <mc:Fallback>
    <p:sp>
      <p:nvSpPr><p:cNvPr id="{{shape_id}}" name="TextBox 1"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
      <p:spPr>
        <a:xfrm><a:off x="{Inches(1)}" y="{Inches(3)}"/><a:ext cx="{Inches(8)}" cy="{Inches(1)}"/></a:xfrm>
        <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
      </p:spPr>
      <p:txBody>
        <a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>FALLBACK PICTURE TEXT</a:t></a:r></a:p>
      </p:txBody>
    </p:sp>
  </mc:Fallback>
</mc:AlternateContent>"""


def image_png(image: SampleImage) -> bytes:
    if image.transparent:
        img = Image.new("RGBA", image.size, (0, 0, 0, 0))
    else:
        img = Image.new("RGB", image.size, "white")
    draw = ImageDraw.Draw(img)
    font = _font()
    for i, line in enumerate(image.lines):
        y = 20 + i * 2.5 * FONT_SIZE
        if draw.textlength(line, font=font) + 40 > image.size[0] or y + FONT_SIZE + 20 > image.size[1]:
            raise ValueError(f"{line!r} does not fit in {image.size}")
        draw.text((20, y), line, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _font() -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    try:
        return ImageFont.truetype(FONT, FONT_SIZE)
    except OSError:
        return ImageFont.load_default(FONT_SIZE)


def make_course_deck(path: Path) -> Path:
    """A 13-slide deck exercising every kind of content; slide numbers are the constants above."""
    prs = Presentation()
    layouts = prs.slide_layouts

    slide = prs.slides.add_slide(layouts[_TITLE_AND_CONTENT])  # BULLETS
    slide.shapes.title.text = "Cell Biology Basics"
    body = slide.placeholders[1].text_frame
    body.text = "Mitochondria produce ATP"
    sub = body.add_paragraph()
    sub.text, sub.level = "Inner membrane folds are cristae", 1
    body.add_paragraph().text = "Ribosomes build proteins"
    body.add_paragraph().text = "Nucleus stores DNA\vand controls the cell"  # \v: a line break
    _add_picture(slide, LOGO, top=6.4, left=7.5, width=2)
    _add_picture(slide, TINY, top=6.4, left=0.2, width=0.2)

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # READING_ORDER, shapes added out of order
    slide.shapes.title.text = "Reading Order"
    _add_text(slide.shapes, "Third box bottom", top=5, left=1)
    _add_text(slide.shapes, "Second box top right", top=2, left=5)
    _add_text(slide.shapes, "First box top left", top=2, left=1)
    _add_picture(slide, LOGO, top=6.4, left=7.5, width=2)

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # GROUPS
    slide.shapes.title.text = "Grouped Shapes"
    _add_text(slide.shapes, "Below the group", top=6, left=1)
    group = slide.shapes.add_group_shape()
    _add_text(group.shapes, "Group inner right", top=2, left=4)
    _add_text(group.shapes, "Group inner left", top=2, left=1)
    nested = group.shapes.add_group_shape()
    _add_text(nested.shapes, "Nested group text", top=3, left=1)
    group.shapes.add_picture(io.BytesIO(image_png(FIGURE)), Inches(1), Inches(4), width=Inches(3))
    _add_picture(slide, LOGO, top=6.4, left=7.5, width=2)

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # TABLE
    slide.shapes.title.text = "Organelles"
    table = slide.shapes.add_table(4, 3, Inches(1), Inches(1.8), Inches(8), Inches(2.5)).table
    rows = (
        ("Organelle", "Function", "Membrane"),
        ("Mitochondrion", "Energy", "Double"),
        ("Lysosome", "Digestion", "Single"),
    )
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    table.cell(3, 0).merge(table.cell(3, 2))
    table.cell(3, 0).text = "Both have membranes"
    _add_picture(slide, FIGURE, top=5, left=1, width=3)

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # CHART
    slide.shapes.title.text = "Enzyme Kinetics"
    data = CategoryChartData()
    data.categories = ["Q1 trial", "Q2 trial", "Q3 trial"]
    data.add_series("Enzyme alpha", (12.5, 17.25, 9.75))
    data.add_series("Enzyme beta", (3.5, 4.25, 8.125))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.8), Inches(4.5), Inches(5), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Reaction rates"
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Rate per second"
    data = CategoryChartData()
    data.categories = ["Liver", "Muscle"]
    data.add_series("Glycogen stores", (100, 400))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.2), Inches(1.8), Inches(4.3), Inches(5), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Storage sites"
    for bar_chart in chart.element.xpath(".//c:barChart"):  # 3-D, which python-pptx cannot author
        bar_chart.tag = qn("c:bar3DChart")

    slide = prs.slides.add_slide(layouts[_TITLE_AND_CONTENT])  # NOTES
    slide.shapes.title.text = "Krebs Cycle"
    slide.placeholders[1].text_frame.text = "Occurs in the matrix"
    slide.notes_slide.notes_text_frame.text = "Remember: the Krebs cycle yields 2 ATP per glucose.\nAsk about NADH."

    slide = prs.slides.add_slide(layouts[_TITLE_AND_CONTENT])  # HIDDEN
    slide.shapes.title.text = "Hidden Bonus Material"
    slide.placeholders[1].text_frame.text = "Only for curious students"
    slide.element.set("show", "0")

    slide = prs.slides.add_slide(layouts[_TITLE_AND_CONTENT])  # FOOTERS
    slide.shapes.title.text = "Footers Are Skipped"
    slide.placeholders[1].text_frame.text = "Body text stays"
    add_footer_placeholders(slide, footer="Confidential footer text", number="Slide 8", date="October 14 2026")

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # PHOTO
    slide.shapes.title.text = "Photosynthesis"
    _add_picture(slide, PHOTO_TEXT, top=1.6, left=1, width=8)
    _add_text(slide.shapes, PHOTO_SLIDE_LINE, top=5.5, left=1)

    slide = prs.slides.add_slide(layouts[_PICTURE_WITH_CAPTION])  # PICTURE_PLACEHOLDER
    slide.shapes.title.text = "Leaf Structure"
    slide.placeholders[1].insert_picture(io.BytesIO(image_png(CAPTIONED)))
    slide.placeholders[2].text_frame.text = "Figure 3: leaf cross-section"

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # SMARTART
    slide.shapes.title.text = "Cell Cycle Phases"
    add_smartart(slide, SMARTART_DATA_XML.encode())

    slide = prs.slides.add_slide(layouts[_TITLE_ONLY])  # EQUATION
    slide.shapes.title.text = "Mass-Energy"
    _add_text(slide.shapes, "Units are joules", top=5, left=1)
    slide.shapes._spTree.append(parse_xml(_EQUATION_XML.format(shape_id=_next_shape_id(slide))))
    _add_text(slide.shapes, "Einstein 1905", top=1.6, left=1)

    prs.slides.add_slide(layouts[_BLANK])  # BLANK

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def add_footer_placeholders(slide: Slide, *, footer: str, number: str, date: str) -> None:
    """Copy the layout's footer, slide-number and date placeholders onto ``slide`` with text.

    python-pptx leaves these out of new slides; PowerPoint adds them when footers are switched on.
    """
    texts = {PP_PLACEHOLDER.FOOTER: footer, PP_PLACEHOLDER.SLIDE_NUMBER: number, PP_PLACEHOLDER.DATE: date}
    for layout_placeholder in slide.slide_layout.placeholders:
        kind = layout_placeholder.placeholder_format.type
        if kind in texts:
            slide.shapes._spTree.append(copy.deepcopy(layout_placeholder.element))
            slide.placeholders[layout_placeholder.placeholder_format.idx].text_frame.text = texts[kind]


def add_smartart(slide: Slide, data_xml: bytes) -> None:
    """Add a SmartArt frame whose data part (``/ppt/diagrams/dataN.xml``) holds ``data_xml``.

    Only the data part is written; the layout, style and colors parts PowerPoint also needs
    are not, so the result is for parsing tests only.
    """
    package = slide.part.package
    partname = package.next_partname("/ppt/diagrams/data%d.xml")
    part = Part(PackURI(partname), CT.DML_DIAGRAM_DATA, package, data_xml)
    rid = slide.part.relate_to(part, RT.DIAGRAM_DATA)
    frame = parse_xml(_SMARTART_FRAME_XML.format(shape_id=_next_shape_id(slide), rid=rid))
    slide.shapes._spTree.append(frame)


def _add_text(shapes, text: str, *, top: float, left: float) -> None:
    shapes.add_textbox(Inches(left), Inches(top), Inches(3), Inches(0.6)).text_frame.text = text


def _add_picture(slide: Slide, image: SampleImage, *, top: float, left: float, width: float) -> None:
    slide.shapes.add_picture(io.BytesIO(image_png(image)), Inches(left), Inches(top), width=Inches(width))


def _next_shape_id(slide: Slide) -> int:
    return max((int(i) for i in slide.element.xpath("//p:cNvPr/@id")), default=0) + 1
