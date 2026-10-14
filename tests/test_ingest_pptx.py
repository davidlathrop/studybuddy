"""Tests for PowerPoint ingestion on decks generated with python-pptx (fast, no network)."""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from PIL.Image import Image

import pptx_samples as samples
from studybuddy.ingest import ocr as ocr_engine
from studybuddy.ingest import pptx as pptx_ingest
from studybuddy.ingest.model import IngestError, PassageDraft


@pytest.fixture(scope="module")
def deck(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return samples.make_course_deck(tmp_path_factory.mktemp("pptx") / "course.pptx")


@pytest.fixture(scope="module")
def drafts(deck: Path) -> list[PassageDraft]:
    return pptx_ingest.extract(deck, ocr=False)


def _text(drafts: list[PassageDraft], slide_no: int, origin: str = "text") -> str:
    (found,) = [d.text for d in drafts if d.slide_no == slide_no and d.origin == origin]
    return found


# --- slide text ----------------------------------------------------------------------------


def test_one_text_passage_per_slide_in_deck_order(drafts: list[PassageDraft]) -> None:
    expected = [(n, "text") for n in range(1, samples.BLANK)]
    expected.insert(samples.NOTES, (samples.NOTES, "notes"))
    assert [(d.slide_no, d.origin) for d in drafts] == expected
    assert all(d.page_no is None and d.t_start_ms is None for d in drafts)


def test_title_comes_first_then_bullets_with_line_breaks(drafts: list[PassageDraft]) -> None:
    assert _text(drafts, samples.BULLETS) == (
        "Cell Biology Basics\nMitochondria produce ATP\nInner membrane folds are cristae\n"
        "Ribosomes build proteins\nNucleus stores DNA\nand controls the cell"
    )


def test_shapes_are_read_top_to_bottom_then_left_to_right(drafts: list[PassageDraft]) -> None:
    assert _text(drafts, samples.READING_ORDER) == (
        "Reading Order\nFirst box top left\nSecond box top right\nThird box bottom"
    )


def test_group_shapes_are_opened_recursively(drafts: list[PassageDraft]) -> None:
    assert _text(drafts, samples.GROUPS) == (
        "Grouped Shapes\nGroup inner left\nGroup inner right\nNested group text\nBelow the group"
    )


def test_table_cells_row_by_row(drafts: list[PassageDraft]) -> None:
    text = _text(drafts, samples.TABLE)
    assert "Organelle | Function | Membrane\nMitochondrion | Energy | Double\nLysosome | Digestion | Single" in text
    assert text.endswith("\nBoth have membranes")  # a merged cell counts once


def test_chart_titles_categories_and_series_but_not_numbers(drafts: list[PassageDraft]) -> None:
    text = _text(drafts, samples.CHART)
    assert text == (
        "Enzyme Kinetics\n"
        "Reaction rates\nRate per second\nQ1 trial, Q2 trial, Q3 trial\nEnzyme alpha, Enzyme beta\n"
        "Storage sites\nLiver, Muscle\nGlycogen stores"  # a 3-D chart, which chart.plots rejects
    )
    assert "12.5" not in text


def test_speaker_notes_are_their_own_passage(drafts: list[PassageDraft]) -> None:
    notes = "Remember: the Krebs cycle yields 2 ATP per glucose.\nAsk about NADH."
    assert _text(drafts, samples.NOTES, "notes") == notes
    assert _text(drafts, samples.NOTES) == "Krebs Cycle\nOccurs in the matrix"


def test_hidden_slide_is_extracted_under_its_position(drafts: list[PassageDraft]) -> None:
    assert _text(drafts, samples.HIDDEN) == "Hidden Bonus Material\nOnly for curious students"
    assert _text(drafts, samples.HIDDEN + 1).startswith("Footers Are Skipped")


def test_footer_slide_number_and_date_placeholders_are_skipped(deck: Path, drafts: list[PassageDraft]) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    slide = Presentation(str(deck)).slides[samples.FOOTERS - 1]
    kinds = {ph.placeholder_format.type for ph in slide.placeholders if ph.has_text_frame and ph.text_frame.text}
    assert {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE} <= kinds  # really present
    assert _text(drafts, samples.FOOTERS) == "Footers Are Skipped\nBody text stays"


def test_picture_placeholder_caption_is_text(drafts: list[PassageDraft]) -> None:
    assert _text(drafts, samples.PICTURE_PLACEHOLDER) == "Leaf Structure\nFigure 3: leaf cross-section"


def test_smartart_text_comes_from_the_diagram_data_part(drafts: list[PassageDraft]) -> None:
    assert _text(drafts, samples.SMARTART) == "Cell Cycle Phases\n" + "\n".join(samples.SMARTART_OUTLINE)


def test_diagram_text_follows_the_outline_not_the_point_list() -> None:
    text = pptx_ingest.diagram_text(samples.SMARTART_DATA_XML.encode())
    assert text.split("\n") == list(samples.SMARTART_OUTLINE)


def test_diagram_text_without_connections_keeps_point_order() -> None:
    xml = samples.SMARTART_DATA_XML.split("<dgm:cxnLst>")[0] + "</dgm:dataModel>"
    in_point_order = ["Mitosis", "S phase", "Interphase", "G2 phase", "G1 phase", "Prophase"]
    assert pptx_ingest.diagram_text(xml.encode()).split("\n") == in_point_order


def test_equation_text_box_hidden_in_alternate_content_is_read(drafts: list[PassageDraft]) -> None:
    text = _text(drafts, samples.EQUATION)
    assert text == f"Mass-Energy\nEinstein 1905\n{samples.EQUATION_TEXT}\nUnits are joules"
    assert "FALLBACK" not in text


def test_empty_slide_has_no_passages(drafts: list[PassageDraft]) -> None:
    assert not [d for d in drafts if d.slide_no == samples.BLANK]


# --- slide count, hidden slides, unhiding --------------------------------------------------


def test_slide_count_and_hidden_slides(deck: Path) -> None:
    assert pptx_ingest.slide_count(deck) == samples.SLIDE_COUNT
    assert pptx_ingest.hidden_slides(deck) == [samples.HIDDEN]


def test_unhide_all_writes_a_visible_copy_and_leaves_the_original(deck: Path, tmp_path: Path) -> None:
    before = hashlib.sha256(deck.read_bytes()).hexdigest()
    copy = tmp_path / "unhidden.pptx"

    assert pptx_ingest.unhide_all(deck, copy) == 1

    assert pptx_ingest.hidden_slides(copy) == []
    assert pptx_ingest.slide_count(copy) == samples.SLIDE_COUNT
    assert pptx_ingest.extract(copy, ocr=False) == pptx_ingest.extract(deck, ocr=False)
    assert hashlib.sha256(deck.read_bytes()).hexdigest() == before
    assert pptx_ingest.hidden_slides(deck) == [samples.HIDDEN]
    assert not list(tmp_path.glob("*.part"))


def test_unhide_all_rewrites_only_the_hidden_slides(deck: Path, tmp_path: Path) -> None:
    copy = tmp_path / "unhidden.pptx"
    pptx_ingest.unhide_all(deck, copy)
    with zipfile.ZipFile(deck) as original, zipfile.ZipFile(copy) as unhidden:
        assert original.namelist() == unhidden.namelist()
        changed = [n for n in original.namelist() if original.read(n) != unhidden.read(n)]
    assert changed == [f"ppt/slides/slide{samples.HIDDEN}.xml"]


def test_unhide_all_without_hidden_slides_returns_zero(deck: Path, tmp_path: Path) -> None:
    once, twice = tmp_path / "once.pptx", tmp_path / "twice.pptx"
    pptx_ingest.unhide_all(deck, once)
    assert pptx_ingest.unhide_all(once, twice) == 0
    assert pptx_ingest.hidden_slides(twice) == []


def test_unhide_all_refuses_to_overwrite_its_source(deck: Path) -> None:
    with pytest.raises(ValueError, match="must differ"):
        pptx_ingest.unhide_all(deck, deck)


# --- OCR -----------------------------------------------------------------------------------


@pytest.fixture
def fake_ocr(monkeypatch: pytest.MonkeyPatch) -> list[tuple[int, int]]:
    """Make OCR available and answer with each sample image's text, chosen by image size."""
    by_size = {image.size: image.text for image in samples.IMAGES}
    calls: list[tuple[int, int]] = []

    def ocr_image(img: Image) -> str:
        calls.append(img.size)
        return by_size[img.size]

    monkeypatch.setattr(ocr_engine, "available", lambda: True)
    monkeypatch.setattr(ocr_engine, "ocr_image", ocr_image)
    return calls


def _ocr_passages(drafts: list[PassageDraft]) -> dict[int, str]:
    return {d.slide_no: d.text for d in drafts if d.origin == "ocr"}


def test_ocr_reads_pictures_and_drops_logos_and_repeated_lines(deck: Path, fake_ocr: list) -> None:
    drafts = pptx_ingest.extract(deck)

    assert _ocr_passages(drafts) == {
        samples.GROUPS: samples.FIGURE.text,  # inside a group; on 2 slides, so kept on both
        samples.TABLE: samples.FIGURE.text,
        samples.PHOTO: "Light reactions in the thylakoid",  # the other line is also slide text
        samples.PICTURE_PLACEHOLDER: samples.CAPTIONED.text,
    }
    assert not any(samples.LOGO.text in d.text for d in drafts)  # on 3 slides: branding
    by_slide = [(d.slide_no, d.origin) for d in drafts if d.slide_no == samples.PHOTO]
    assert by_slide == [(samples.PHOTO, "text"), (samples.PHOTO, "ocr")]


def test_each_distinct_image_is_read_once_and_tiny_ones_never(deck: Path, fake_ocr: list) -> None:
    pptx_ingest.extract(deck)
    assert sorted(fake_ocr) == sorted(image.size for image in samples.IMAGES if image is not samples.TINY)


def test_ocr_false_never_reads_pictures(deck: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def fail(img: Image) -> str:
        raise AssertionError("OCR must not run")

    monkeypatch.setattr(ocr_engine, "available", lambda: True)
    monkeypatch.setattr(ocr_engine, "ocr_image", fail)
    assert not _ocr_passages(pptx_ingest.extract(deck, ocr=False))


def test_ocr_unavailable_gives_no_ocr_passages(deck: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ocr_engine, "available", lambda: False)
    assert pptx_ingest.extract(deck) == pptx_ingest.extract(deck, ocr=False)


@pytest.mark.skipif(not ocr_engine.available(), reason="needs macOS Vision OCR (ocrmac)")
def test_real_ocr_on_slide_pictures(deck: Path) -> None:
    found = {n: text.casefold() for n, text in _ocr_passages(pptx_ingest.extract(deck)).items()}

    assert set(found) == {samples.GROUPS, samples.TABLE, samples.PHOTO, samples.PICTURE_PLACEHOLDER}
    assert "thylakoid" in found[samples.PHOTO] and "calvin" not in found[samples.PHOTO]
    assert "stomata" in found[samples.PICTURE_PLACEHOLDER]  # transparent PNG
    assert "membrane transport" in found[samples.GROUPS]
    assert not any("acme" in text for text in found.values())


# --- bad input -----------------------------------------------------------------------------


def _truncated(deck: Path, path: Path) -> None:
    data = deck.read_bytes()
    path.write_bytes(data[: len(data) // 2])


def _zip_without_content_types(deck: Path, path: Path) -> None:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("hello.txt", "not a presentation")


def _word_document(deck: Path, path: Path) -> None:
    docx = pytest.importorskip("docx")
    docx.Document().save(str(path))


@pytest.mark.parametrize(
    ("make", "message"),
    [
        (lambda deck, path: path.write_text("just some text"), "not a readable .pptx"),
        (_truncated, "not a readable .pptx"),
        (_zip_without_content_types, "not a readable .pptx"),
        (_word_document, "not a readable .pptx"),
        (lambda deck, path: path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504)), "legacy .ppt"),
    ],
    ids=["text", "truncated", "zip", "docx", "ole"],
)
def test_unreadable_files_raise_ingest_error(deck: Path, tmp_path: Path, make, message: str) -> None:
    bad = tmp_path / "bad.pptx"
    make(deck, bad)
    for call in (pptx_ingest.extract, pptx_ingest.slide_count, pptx_ingest.hidden_slides):
        with pytest.raises(IngestError, match=message):
            call(bad)
    with pytest.raises(IngestError, match=message):
        pptx_ingest.unhide_all(bad, tmp_path / "copy.pptx")


def test_missing_file_raises_ingest_error(tmp_path: Path) -> None:
    with pytest.raises(IngestError, match="no such file"):
        pptx_ingest.extract(tmp_path / "missing.pptx")


def test_a_broken_shape_is_skipped_not_fatal(deck: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def broken(chart) -> str:
        raise KeyError("rId9")

    monkeypatch.setattr(pptx_ingest, "_chart_text", broken)
    drafts = pptx_ingest.extract(deck, ocr=False)
    assert _text(drafts, samples.CHART) == "Enzyme Kinetics"
    assert _text(drafts, samples.TABLE).startswith("Organelles\nOrganelle | Function")
